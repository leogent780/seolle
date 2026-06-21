from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

def set_table_style(table):
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)

def add_heading(doc, text, level=1, color=None):
    h = doc.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in h.runs:
        if color:
            run.font.color.rgb = RGBColor(*color)
    return h

def add_table(doc, headers, rows, header_bg=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.size = Pt(10)
        if header_bg:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), header_bg)
            tcPr.append(shd)
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = str(val)
            cell.paragraphs[0].runs[0].font.size = Pt(10)
    return table

# ─────────────────────────────────────────────
# 고객공유용
# ─────────────────────────────────────────────
doc1 = Document()
doc1.core_properties.title = 'Seoul Glow Day — 당일 변신 패키지'

style = doc1.styles['Normal']
style.font.name = 'Malgun Gothic'
style.font.size = Pt(10)

add_heading(doc1, '✨ Seoul Glow Day — 당일 변신 패키지', 1, (31, 73, 125))
doc1.add_paragraph('하루 안에 피부·외모 변화를 경험하는 K-뷰티 데이트립\n명동 전 구간 도보 이동 | 소그룹 운영').italic = True

doc1.add_paragraph()
add_heading(doc1, '💰 패키지 가격', 2)
add_table(doc1,
    ['구분', '가격'],
    [['기본 패키지', '₩399,000'], ['큐레이터 동행', '₩50,000~ 별도']],
    'D9E2F3'
)

doc1.add_paragraph()
add_heading(doc1, '📦 기본 패키지에 포함된 것', 2)
add_table(doc1,
    ['항목', '내용'],
    [
        ['퍼스널컬러 분석', '웜/쿨톤 · 계절 유형 · 얼굴형 · 체형 분석'],
        ['헤어 & 메이크업', '명동 제휴 헤어숍 스타일링 + 한국식 글로우 메이크업'],
        ['프로필 촬영', '실내 스튜디오 · 전문 포토그래퍼'],
        ['한국 eSIM', '여행 기간 데이터 사용'],
        ['T-money 교통카드', '대중교통 이용 가능'],
        ['여행자 보험', '당일 여행 중 보장'],
        ['업체별 예약 확인서', '코스별 예약 정보 사전 발송'],
        ['카카오톡 긴급 연락', '한국어 · 영어 응대'],
    ],
    'D9E2F3'
)

doc1.add_paragraph()
add_heading(doc1, '🗓️ 하루 스케줄', 2)
add_table(doc1,
    ['시간', '코스', '비고'],
    [
        ['10:00', '퍼스널컬러 분석', '포함'],
        ['11:00', '올리브영 & 약국 쇼핑', '쇼핑비 개별 결제'],
        ['13:30', '피부진단 & 에스테틱', '시술비 개별 결제'],
        ['15:30', '헤어 & 메이크업', '포함'],
        ['17:30', '프로필 촬영', '포함'],
        ['19:30', '해산', '명동역 인근'],
    ],
    'D9E2F3'
)
doc1.add_paragraph('※ 교통비는 개별 결제 (T-money 카드 제공)').italic = True

doc1.add_paragraph()
add_heading(doc1, '🏥 병원 시술 예약 (별도)', 2)
p = doc1.add_paragraph('당일 코스와 별개로 ')
p.add_run('다른 날짜').bold = True
p.add_run('에 명동 TOP5 피부과·성형외과 시술 예약을 도와드립니다.')
for item in ['원하는 시술 + 날짜 선택', '영어 · 중국어 · 일본어 상담 가능 병원 연결', '예약 확인서 발송', '시술비는 병원 현장 결제']:
    doc1.add_paragraph(item, style='List Bullet')

doc1.add_paragraph()
add_heading(doc1, '💆 피부진단 & 에스테틱 시술 메뉴 (현장 결제)', 2)
add_table(doc1,
    ['시술', '소요 시간', '가격'],
    [
        ['기초 페이셜 클렌징', '60분', '₩50,000–80,000'],
        ['수분 앰플 집중 관리', '60분', '₩80,000–120,000'],
        ['블랙헤드 & 모공 관리', '60분', '₩80,000–100,000'],
    ],
    'D9E2F3'
)

doc1.add_paragraph()
add_heading(doc1, '👩‍💼 큐레이터 동행 옵션', 2)
doc1.add_paragraph('전담 큐레이터가 원하는 구간만 함께 이동하며 통역 · 쇼핑 가이드 · 병원 동행을 도와드립니다.')
doc1.add_paragraph('※ 구간은 반드시 연속으로 선택 (중간 건너뛰기 불가)').italic = True
add_table(doc1,
    ['동행 구간', '추가 요금'],
    [
        ['1구간', '₩50,000'],
        ['2구간', '₩80,000'],
        ['3구간', '₩110,000'],
        ['4구간', '₩130,000'],
        ['5구간 (전 구간)', '₩150,000'],
    ],
    'D9E2F3'
)
doc1.add_paragraph()
doc1.add_paragraph('큐레이터 제공 서비스', style='List Bullet').runs[0].bold = True
for item in ['영어 · 중국어 · 일본어 통역', '퍼스널컬러 기반 쇼핑 1:1 추천', '피부과 · 에스테틱 상담 동행']:
    doc1.add_paragraph(item, style='List Bullet')

doc1.add_paragraph()
add_heading(doc1, '❌ 포함되지 않는 것', 2)
for item in ['교통비 (지하철 · 택시)', '점심식사', '올리브영 · 약국 쇼핑비', '에스테틱 시술비', '병원 시술비']:
    doc1.add_paragraph(item, style='List Bullet')

doc1.add_paragraph()
add_heading(doc1, '📌 예약 안내', 2)
for item in ['최소 3일 전 예약 필요', '카카오톡 채널 또는 이메일로 문의', '병원 시술 예약은 5일 전 요청 권장']:
    doc1.add_paragraph(item, style='List Bullet')

doc1.save('/home/user/seolle/packages/고객공유용_당일변신패키지.docx')
print('고객공유용 완료')

# ─────────────────────────────────────────────
# 내부공유용
# ─────────────────────────────────────────────
doc2 = Document()
doc2.core_properties.title = '[내부용] Seoul Glow Day 원가 & 수익 구조'

style2 = doc2.styles['Normal']
style2.font.name = 'Malgun Gothic'
style2.font.size = Pt(10)

add_heading(doc2, '[내부용] Seoul Glow Day — 원가 & 수익 구조', 1, (192, 0, 0))
doc2.add_paragraph('🔒 내부 공유 전용 — 외부 배포 금지').runs[0].bold = True

doc2.add_paragraph()
add_heading(doc2, '[1] 패키지 코스 구성', 2)
add_table(doc2,
    ['시간', '코스', '포함 여부', '결제 주체'],
    [
        ['10:00', '퍼스널컬러 분석', '✅ 포함', '우리'],
        ['11:00', '올리브영 & 약국 쇼핑', '❌ 미포함', '고객 현장'],
        ['13:30', '피부진단 & 에스테틱', '❌ 미포함', '고객 현장'],
        ['15:30', '헤어 & 메이크업', '✅ 포함', '우리'],
        ['17:30', '프로필 촬영', '✅ 포함', '우리'],
        ['19:30', '해산', '—', '—'],
        ['별도 날짜', '명동 TOP5 병원 예약 대행', '예약만 우리', '고객 현장'],
    ],
    'FCE4D6'
)

doc2.add_paragraph()
add_heading(doc2, '[2] 구성 상품별 원가', 2)
add_heading(doc2, '실체 상품 (핵심 3개)', 3)
add_table(doc2,
    ['항목', 'B2B 원가', '비고'],
    [
        ['퍼스널컬러 분석', '₩60,000', '예상 원가 (협약 전)'],
        ['헤어 & 메이크업', '₩70,000', '협약 전 예상가'],
        ['프로필 촬영', '₩50,000', '협약 완료'],
        ['소계', '₩180,000', ''],
    ],
    'FCE4D6'
)

doc2.add_paragraph()
add_heading(doc2, '번들 상품 (원가율 낮추기용)', 3)
add_table(doc2,
    ['항목', '원가', '고객 인식 가치', '원가율'],
    [
        ['여행자 보험', '₩2,500', '₩12,000', '21%'],
        ['eSIM', '₩10,000', '₩30,000', '33%'],
        ['T-money 교통카드', '₩2,500', '₩10,000', '25%'],
        ['카카오톡 긴급 연락', '₩0', '₩5,000', '0%'],
        ['면세 쿠폰', '₩0', '₩5,000', '0%'],
        ['예약 확인서 발송', '₩0', '₩3,000', '0%'],
        ['소계', '₩15,000', '₩65,000', '23%'],
    ],
    'FCE4D6'
)

doc2.add_paragraph()
add_heading(doc2, '[3] 원가 합계', 2)
add_table(doc2,
    ['구분', '금액'],
    [
        ['실체 상품', '₩180,000'],
        ['번들 상품', '₩15,000'],
        ['운영비', '₩15,000'],
        ['총 원가', '₩210,000'],
    ],
    'FCE4D6'
)

doc2.add_paragraph()
add_heading(doc2, '[4] 패키지 판매가 & 원가율', 2)
add_table(doc2,
    ['판매가', '원가', '원가율', '마진'],
    [
        ['₩350,000', '₩210,000', '60.0%', '₩140,000'],
        ['₩380,000', '₩210,000', '55.3%', '₩170,000'],
        ['₩399,000 ✅ 적용', '₩210,000', '52.6%', '₩189,000'],
    ],
    'FCE4D6'
)

doc2.add_paragraph()
add_heading(doc2, '[5] 큐레이터 수익 구조 (중개 수수료 20%)', 2)
add_table(doc2,
    ['구간', '고객 지불', '큐레이터 지급 (80%)', '우리 수익 (20%)'],
    [
        ['1구간', '₩50,000', '₩40,000', '₩10,000'],
        ['2구간', '₩80,000', '₩64,000', '₩16,000'],
        ['3구간', '₩110,000', '₩88,000', '₩22,000'],
        ['4구간', '₩130,000', '₩104,000', '₩26,000'],
        ['5구간', '₩150,000', '₩120,000', '₩30,000'],
    ],
    'FCE4D6'
)

doc2.add_paragraph()
add_heading(doc2, '[6] 큐레이터 포함 시 총 수익 & 전체 원가율 (기본 ₩399,000 기준)', 2)
add_table(doc2,
    ['큐레이터 구간', '총 판매가', '총 원가', '전체 원가율', '총 마진'],
    [
        ['없음', '₩399,000', '₩210,000', '52.6%', '₩189,000'],
        ['1구간', '₩449,000', '₩250,000', '55.7%', '₩199,000'],
        ['2구간', '₩479,000', '₩274,000', '57.2%', '₩205,000'],
        ['3구간', '₩509,000', '₩298,000', '58.5%', '₩211,000'],
        ['4구간', '₩529,000', '₩314,000', '59.4%', '₩215,000'],
        ['5구간', '₩549,000', '₩330,000', '60.1% ⚠️', '₩219,000'],
    ],
    'FCE4D6'
)

doc2.add_paragraph()
add_heading(doc2, '[7] 현장결제 항목 수익 검토', 2)
add_table(doc2,
    ['항목', '고객 결제', '우리 수익'],
    [
        ['올리브영 · 약국 쇼핑', '고객 직접', '면세쿠폰 리베이트 가능성 검토'],
        ['피부진단 & 에스테틱', '고객 직접', '예약 대행 수수료 협약 검토'],
        ['명동 TOP5 병원 시술', '고객 직접', '예약 대행 수수료 협약 검토'],
    ],
    'FCE4D6'
)

doc2.add_paragraph()
add_heading(doc2, '[8] 협약 미완료 항목', 2)
add_table(doc2,
    ['항목', '상태', '비고'],
    [
        ['퍼스널컬러 전문가', '🔴 미완료', '원가 ₩60,000 예상'],
        ['헤어 & 메이크업', '🔴 미완료', '원가 ₩70,000 예상'],
        ['프로필 촬영', '🟢 완료', '₩50,000 확정'],
        ['에스테틱 업체', '🔴 미완료', '예약 대행 수수료 협의 필요'],
        ['명동 TOP5 병원', '🔴 미완료', '협약 + 외국인 특별가 확인'],
        ['eSIM 공급처', '🔴 미완료', '도매 or 대리점 제휴'],
        ['여행자 보험사', '🔴 미완료', '대행 수수료 구조 확인'],
    ],
    'FCE4D6'
)

doc2.save('/home/user/seolle/packages/내부공유용_당일변신패키지.docx')
print('내부공유용 완료')
