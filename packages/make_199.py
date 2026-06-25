from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def cell_text(cell, lines, bold_first=False, font_size=9):
    cell.paragraphs[0].clear()
    for i, line in enumerate(lines):
        if i == 0:
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
        run = p.add_run(line)
        run.font.size = Pt(font_size)
        run.font.name = 'Malgun Gothic'
        if i == 0 and bold_first:
            run.bold = True
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)

def add_heading(doc, text, level=1, color=None):
    h = doc.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in h.runs:
        run.font.name = 'Malgun Gothic'
        if color:
            run.font.color.rgb = RGBColor(*color)
    return h

def add_simple_table(doc, headers, rows, header_bg='D9E2F3'):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(10)
        cell.paragraphs[0].runs[0].font.name = 'Malgun Gothic'
        set_cell_bg(cell, header_bg)
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].text = str(val)
            row.cells[c_idx].paragraphs[0].runs[0].font.size = Pt(10)
            row.cells[c_idx].paragraphs[0].runs[0].font.name = 'Malgun Gothic'
    return table

def add_schedule_row(table, row_idx, time_text, course_lines, note_text):
    row = table.rows[row_idx]
    row.cells[0].text = time_text
    for p in row.cells[0].paragraphs:
        for r in p.runs:
            r.font.size = Pt(9)
            r.font.name = 'Malgun Gothic'
    cell_text(row.cells[1], course_lines, bold_first=True, font_size=9)
    row.cells[2].text = note_text
    for p in row.cells[2].paragraphs:
        for r in p.runs:
            r.font.size = Pt(9)
            r.font.name = 'Malgun Gothic'

def add_day_table(doc, day_data):
    table = doc.add_table(rows=1 + len(day_data), cols=3)
    table.style = 'Table Grid'
    for i, h in enumerate(['시간', '코스', '안내']):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(10)
        cell.paragraphs[0].runs[0].font.name = 'Malgun Gothic'
        set_cell_bg(cell, 'D9E2F3')
    for i, (t, c, n) in enumerate(day_data):
        add_schedule_row(table, i + 1, t, c, n)
    for row in table.rows:
        row.cells[0].width = Cm(2.0)
        row.cells[1].width = Cm(11.5)
        row.cells[2].width = Cm(2.5)
    return table

# ─────────────────────────────────────────────
doc = Document()
doc.core_properties.title = 'Seoul Renewal — 2박 3일 토탈 뷰티 리뉴얼'

style = doc.styles['Normal']
style.font.name = 'Malgun Gothic'
style.font.size = Pt(10)

# 제목
add_heading(doc, '✨ Seoul Renewal — 2박 3일 토탈 뷰티 리뉴얼', 1, (31, 73, 125))
p = doc.add_paragraph('피부·반영구·웰니스까지 — 한국에서 완전히 새로워지는 3일')
p.runs[0].italic = True
p.runs[0].font.name = 'Malgun Gothic'
p2 = doc.add_paragraph('명동·강남 일대  |  큐레이터 동행 옵션')
p2.runs[0].italic = True
p2.runs[0].font.name = 'Malgun Gothic'

# ── DAY 1 ──
doc.add_paragraph()
add_heading(doc, 'DAY 1 — 진단 & 피부 정돈', 2, (31, 73, 125))

day1_data = [
    (
        '10:00\n~\n11:30',
        [
            '퍼스널컬러 분석 1:1  (90분)',
            '· 웜톤/쿨톤 세부 계절 유형 분류 (봄·여름·가을·겨울)',
            '· 얼굴형 & 체형 분석 / 통역 포함',
            '',
            '추천 업체 TOP5',
            '① 로아 퍼스널컬러  ② 퍼스널컬러랩  ③ 뮤즈클로젯',
            '④ 위드컬러  ⑤ 마이컬러스튜디오',
            '',
            '💡 3일 전체 스타일링·시술의 방향을 결정하는 첫 단계.',
            '   이 결과 없이는 이후 모든 코스가 흐릿해짐',
        ],
        '✅ 포함'
    ),
    (
        '11:30\n~\n13:30',
        [
            '올리브영 & 약국 쇼핑  (2시간)',
            '· 퍼스널컬러 결과 직후 바로 쇼핑 적용',
            '· 명동 올리브영 메가스토어 + 인근 약국',
            '',
            '💡 진단 결과가 가장 선명할 때 바로 쇼핑',
        ],
        '💳 쇼핑비\n개별 결제'
    ),
    (
        '13:30\n~\n15:00',
        [
            '네일 케어  (90분)',
            '· 손톱 & 발톱 기본 관리',
            '· 젤 네일 컬러링 (퍼스널컬러 기반 추천)',
            '· 큐티클 정리 + 핸드케어',
            '',
            '추천 업체 TOP5',
            '① 네일팝  ② 에뛰드 네일  ③ 글로시네일',
            '④ 핑크네일 명동  ⑤ 네일아트 클럽',
            '',
            '💡 여행 첫날 네일을 완성하면 남은 일정 내내 손이 예뻐 보임.',
            '   쇼핑 직후 퍼스널컬러 기반 컬러 선택으로 구매한 제품과 룩 통일',
        ],
        '✅ 포함'
    ),
    (
        '15:00\n~\n16:00',
        [
            '티 오마카세  (1시간)',
            '· 전통차 소믈리에가 선별한 계절 차 코스 제공',
            '· 차 설명 + 다과 페어링 포함',
            '· 명동·인사동 인근 프리미엄 티룸',
            '',
            '추천 업체 TOP5',
            '① 오설록 티하우스  ② 차마시는 뜰  ③ 티컬렉티브',
            '④ 명가원  ⑤ 오우재',
            '',
            '💡 오후 시술 전 충분한 휴식.',
            '   프리미엄 차 코스로 패키지 고급감 UP',
        ],
        '✅ 포함'
    ),
    (
        '16:00\n~\n18:00',
        [
            '피부진단 & 에스테틱  (2시간)',
            '· 전문 에스테티션 1:1 피부 진단 (20분)',
            '· 시술 선택 (60–80분)',
            '  — 트러블집중관리 + 진정관리',
            '  — 얼굴윤곽 + 탄력관리',
            '  — 화이트닝 + 탄력관리',
            '',
            '추천 업체 TOP5',
            '① 오슬 스킨케어  ② 라뷰 에스테틱  ③ 더스킨랩',
            '④ 뷰티풀스킨  ⑤ 스킨케어 클리닉',
            '',
            '💡 반영구 시술 전날 미리 피부 정돈.',
            '   피부가 안정된 상태에서 다음 날 시술 효과 극대화',
        ],
        '✅ 포함'
    ),
    (
        '18:00\n~\n19:30',
        [
            '크라이오 + 산소캡슐 웰니스  (90분)',
            '· 크라이오테라피 — 냉기 자극으로 혈액순환·피부 탄력',
            '· 산소캡슐 — 고압 산소로 피부 재생 & 피로 회복',
            '',
            '추천 업체 TOP5',
            '① 크라이오서울  ② 바디럽  ③ 웰니스클럽 강남',
            '④ 리커버리랩  ⑤ 아이스에이지 웰니스',
            '',
            '💡 에스테틱 직후 웰니스 처치로 피부 진정·회복 극대화',
        ],
        '✅ 포함'
    ),
    (
        '19:30',
        ['DAY 1 해산', '· 명동역 인근'],
        '—'
    ),
]

add_day_table(doc, day1_data)

# ── DAY 2 ──
doc.add_paragraph()
add_heading(doc, 'DAY 2 — 반영구 & 스페셜 케어', 2, (31, 73, 125))

day2_data = [
    (
        '10:00\n~\n12:30',
        [
            '반영구 시술 1종 선택  (2–2.5시간)',
            '· 눈썹 문신 — 자연스러운 결 눈썹',
            '· 립 문신 — 혈색 있는 핑크/코럴 립',
            '· 아이라인 문신 — 눈매 또렷하게',
            '· 두피 문신 — 헤어라인 & 볼륨 보정',
            '',
            '추천 업체 TOP5',
            '① 오브뷰티  ② 미소반영구  ③ 아트메이크업랩',
            '④ 뷰티펙트  ⑤ 라인아티스트',
            '',
            '💡 한 번 시술로 1–2년 유지.',
            '   퍼스널컬러 기반 색조 선택으로 얼굴 전체 밸런스 완성',
        ],
        '✅ 포함'
    ),
    (
        '13:00\n~\n14:00',
        [
            '한의원 약침 & 시술  (1시간)',
            '· 피부 약침 — 콜라겐 재생·탄력',
            '· 혈자리 침 — 얼굴 윤곽·리프팅',
            '· 한방 팩 마무리',
            '',
            '추천 업체 TOP5',
            '① 강남 경희한의원  ② 명동 자생한의원  ③ 아이 한의원',
            '④ 피부한의원 명동  ⑤ 동의보감한의원',
            '',
            '💡 반영구 시술 후 내부에서 피부를 강화.',
            '   한방 약침은 빠른 회복 + 글로우 효과',
        ],
        '✅ 포함'
    ),
    (
        '14:00\n~\n15:30',
        [
            '헤드스파 & 두피 케어  (90분)',
            '· 두피 진단 (두피 타입·탈모 여부)',
            '· 한방 헤드스파 트리트먼트',
            '· 두피 스케일링 + 영양 앰플',
            '',
            '추천 업체 TOP5',
            '① 모스트헤어  ② 헤드스파 클리닉  ③ 샴푸헤어',
            '④ 두피케어랩  ⑤ 스칼프 클리닉 강남',
            '',
            '💡 두피가 건강해야 모발이 살아남.',
            '   두피문신 선택 시 시술 전 두피 상태 점검으로 활용 가능',
        ],
        '✅ 포함'
    ),
    (
        '15:30\n~\n16:30',
        [
            '치아미백  (1시간)',
            '· 전문 치아미백 시술 (레이저 or LED)',
            '· 시술 전후 색상 비교',
            '',
            '추천 업체 TOP5',
            '① 유디치과  ② 스마일라인 치과  ③ 이룸치과',
            '④ 명동 뷰티덴탈  ⑤ 화이트클리닉',
            '',
            '💡 미소가 바뀌면 전체 인상이 달라짐.',
            '   촬영 전날 시술해 최상의 상태로 마무리',
        ],
        '✅ 포함'
    ),
    (
        '17:00\n~\n18:30',
        [
            '전신 마사지  (90분)',
            '· 아로마 오일 전신 마사지',
            '· 2일간 시술 피로 회복',
            '· 혈액순환·부기 완화',
            '',
            '추천 업체 TOP5',
            '① 힐링스파 명동  ② 바디웍스  ③ 아로마힐링 강남',
            '④ 스파위드  ⑤ 웰니스 스파 명동',
            '',
            '💡 2일간 집중 시술로 쌓인 피로를 해소.',
            '   마지막 날 촬영을 위해 몸과 마음 리셋',
        ],
        '✅ 포함'
    ),
    (
        '19:00',
        ['DAY 2 해산', '· 명동역 인근'],
        '—'
    ),
]

add_day_table(doc, day2_data)

# ── DAY 3 ──
doc.add_paragraph()
add_heading(doc, 'DAY 3 — 피니싱 & 촬영', 2, (31, 73, 125))

day3_data = [
    (
        '10:00\n~\n11:00',
        [
            '헤어 & 메이크업  (1시간)  ★ 옵션',
            '· 한국식 데일리 헤어 스타일링',
            '· 퍼스널컬러 기반 내추럴 글로우 메이크업',
            '',
            '추천 업체 TOP5',
            '① 박준뷰티랩  ② 준오헤어  ③ 이철헤어커커',
            '④ 토니앤가이  ⑤ 헤라뷰티스튜디오',
            '',
            '💡 촬영 직전 마무리 스타일링.',
            '   기본 패키지에는 미포함 — 원하는 분만 선택',
        ],
        '🔶 옵션\n+₩150,000'
    ),
    (
        '11:00\n~\n13:00',
        [
            '프로필 촬영  (2시간)',
            '· 명동 인근 실내 스튜디오',
            '· 전문 포토그래퍼',
            '· 데일리 룩 + 글로우 컨셉',
            '',
            '추천 업체 TOP5',
            '① 필름드 스튜디오  ② 모멘트 포토  ③ 셀프리 스튜디오',
            '④ 데일리 스냅  ⑤ 피치 스튜디오',
            '',
            '💡 3일간의 변화가 모두 담긴 최상의 상태.',
            '   반영구·피부·네일 모두 완성된 후 촬영',
        ],
        '✅ 포함'
    ),
    (
        '13:00',
        ['DAY 3 해산', '· 명동역 인근'],
        '—'
    ),
    (
        '별도\n날짜',
        [
            '명동 TOP5 병원 시술 예약 대행',
            '· 퍼스널컬러 진단 결과 기반 시술 추천',
            '· 피부과 / 성형외과 중 선택',
            '· 시술비는 병원 현장 결제',
        ],
        '🏥 예약\n대행'
    ),
]

add_day_table(doc, day3_data)

# ── 포함 항목 ──
doc.add_paragraph()
add_heading(doc, '🎁 포함 항목', 2)
add_heading(doc, 'Seoul Renewal 패키지 포함', 3)
for item in [
    'DAY1  퍼스널컬러 분석 1:1 (통역포함)',
    'DAY1  네일 케어',
    'DAY1  피부진단 & 에스테틱',
    'DAY1  티 오마카세 (전통차 코스 + 다과)',
    'DAY1  크라이오 + 산소캡슐 웰니스',
    'DAY2  반영구 시술 1종 선택 (눈썹/립/아이라인/두피)',
    'DAY2  한의원 약침 & 시술',
    'DAY2  헤드스파 & 두피 케어',
    'DAY2  치아미백',
    'DAY2  전신 마사지',
    'DAY3  프로필 촬영',
]:
    p = doc.add_paragraph(item, style='List Bullet')
    p.runs[0].font.name = 'Malgun Gothic'

add_heading(doc, '여행 편의 키트', 3)
for item in [
    '한국 eSIM',
    '여행자 보험',
    '코스별 예약 확인서',
    '카카오톡 긴급 연락',
    '면세 쇼핑 쿠폰',
    '웰컴 K뷰티 샘플팩',
]:
    p = doc.add_paragraph(item, style='List Bullet')
    p.runs[0].font.name = 'Malgun Gothic'

add_heading(doc, '★ 별도 옵션', 3)
for item in [
    '헤어 & 메이크업 (+₩150,000)',
]:
    p = doc.add_paragraph(item, style='List Bullet')
    p.runs[0].font.name = 'Malgun Gothic'

# ── 가격 안내 ──
doc.add_paragraph()
add_heading(doc, '💰 가격 안내', 2)
p = doc.add_paragraph()
r = p.add_run('기본 패키지  ₩1,990,000')
r.bold = True
r.font.size = Pt(12)
r.font.name = 'Malgun Gothic'

doc.add_paragraph()
add_simple_table(doc,
    ['큐레이터 동행 구간', '추가 요금', '총 금액'],
    [
        ['없음', '—', '₩1,990,000'],
        ['1구간', '+₩50,000', '₩2,040,000'],
        ['2구간', '+₩80,000', '₩2,070,000'],
        ['3구간', '+₩110,000', '₩2,100,000'],
        ['4구간', '+₩130,000', '₩2,120,000'],
        ['5구간 (전 구간)', '+₩150,000', '₩2,140,000'],
    ],
    'D9E2F3'
)
doc.add_paragraph('※ 큐레이터는 원하는 구간만 연속으로 선택 가능합니다').runs[0].font.name = 'Malgun Gothic'
doc.add_paragraph('※ 지원 언어: 영어 · 중국어 · 일본어').runs[0].font.name = 'Malgun Gothic'

doc.add_paragraph()
add_simple_table(doc,
    ['항목', '결제'],
    [
        ['명동 병원 시술비', '💳 현장결제 (별도 날짜)'],
        ['쇼핑비', '💳 개별 결제'],
        ['헤어 & 메이크업 (옵션)', '💳 +₩150,000'],
    ],
    'FFF2CC'
)

# ── 원가 구조 (내부) ──
doc.add_paragraph()
add_heading(doc, '📊 원가 구조 (내부 전용)', 2, (192, 0, 0))
add_simple_table(doc,
    ['항목', '원가'],
    [
        ['퍼스널컬러 분석 1:1 (통역포함)', '₩90,000'],
        ['피부진단 & 에스테틱', '₩120,000'],
        ['티 오마카세', '₩40,000'],
        ['크라이오 + 산소캡슐', '₩100,000'],
        ['반영구 시술 1종', '₩280,000'],
        ['한의원 약침', '₩80,000'],
        ['헤드스파 & 두피케어', '₩70,000'],
        ['치아미백', '₩120,000'],
        ['전신 마사지', '₩80,000'],
        ['네일 케어', '₩60,000'],
        ['프로필 촬영', '₩130,000'],
        ['번들 (eSIM·보험 등)', '₩15,000'],
        ['운영비', '₩30,000'],
        ['원가 합계', '₩1,215,000'],
        ['원가율', '61.1% ⚠️'],
    ],
    'FCE4D6'
)

out = '/home/user/seolle/packages/고객공유용_199만원_SeoulRenewal.docx'
doc.save(out)
print('완료:', out)
